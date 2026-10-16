"""Word / Excel report export (nma_report): the .docx is a valid OOXML
package, the .xlsx has the expected sheets, figures are PNGs and the
optional sections (Bayesian, inconsistency, CINeMA, contribution,
absolute effects) appear when their inputs are given."""

import io
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import nma_editor_support as sup  # noqa: E402
import nma_report  # noqa: E402
from heor_tree_engine.nma_absolute import absolute_effects  # noqa: E402
from heor_tree_engine.nma_cinema import run_cinema  # noqa: E402
from heor_tree_engine.nma_contribution import contribution_matrix  # noqa: E402
from heor_tree_engine.nma_freq import run_nma_frequentist  # noqa: E402
from heor_tree_engine.nma_inconsistency import run_inconsistency  # noqa: E402
from heor_tree_engine.nma_network import summarize_network  # noqa: E402
from heor_tree_engine.nma_parser import resolve_nma  # noqa: E402

import openpyxl  # noqa: E402
from functools import lru_cache  # noqa: E402


@lru_cache(maxsize=1)
def basic():
    document = sup.default_document()
    model = resolve_nma(document)
    result = run_nma_frequentist(model)
    assert result.errors == []
    return document, model, summarize_network(model), result


def _docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = set(zf.namelist())
        assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml",
                "word/_rels/document.xml.rels"} <= names
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                ElementTree.fromstring(zf.read(name))  # well-formed XML
        media = [n for n in names if n.startswith("word/media/")]
        assert len(media) >= 2  # network + forest
        for name in media:
            assert zf.read(name)[:8] == b"\x89PNG\r\n\x1a\n"
        root = ElementTree.fromstring(zf.read("word/document.xml"))
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return " ".join(t.text or "" for t in root.iter(f"{ns}t"))


def test_frequentist_report_is_valid_docx_and_xlsx() -> None:
    document, model, summary, result = basic()
    bundle = nma_report.build_report(nma_report.ReportInput(
        document=document, model=model, summary=summary, result=result, title="Smoke report"))

    assert bundle.filename_stem.startswith("smoke_report")
    assert set(bundle.figures) == {"network", "forest"}
    text = _docx_text(bundle.docx_bytes)
    for label in ("Methods", "League table", "Ranking", "Smoke report"):
        assert label in text
    assert "Bayesian" not in text.split("Methods")[0]

    wb = openpyxl.load_workbook(io.BytesIO(bundle.xlsx_bytes))
    for sheet in ("Effects", "League", "Ranking", "Data arms"):
        assert any(sheet.lower() in name.lower() for name in wb.sheetnames), wb.sheetnames
    assert wb.sheetnames[0] == "Report" and wb["Report"]["A1"].value == "Smoke report"
    effects = wb["Effects"]
    assert effects.max_row >= 2 and effects.max_column >= 3


def test_optional_sections_are_included_and_translated() -> None:
    document, model, summary, result = basic()
    inconsistency = run_inconsistency(model)
    contribution = contribution_matrix(model, effects="random", nma=result)
    cinema = run_cinema(model, result, contribution=contribution, inconsistency=inconsistency)
    absolute = absolute_effects(result.random, result.measure, 0.2)
    labels = {"sec_methods": "Methods (custom)", "sec_league": "League table (custom)", "sec_cinema": "CINeMA"}
    bundle = nma_report.build_report(nma_report.ReportInput(
        document=document, model=model, summary=summary, result=result, labels=labels, language="en",
        inconsistency=inconsistency, cinema=cinema, contribution=contribution, absolute=absolute,
        baseline_risk=0.2, title="NMA report", author="Trang Nguyen"))

    text = _docx_text(bundle.docx_bytes)
    assert "Methods (custom)" in text and "League table (custom)" in text and "CINeMA" in text and "Trang Nguyen" in text
    wb = openpyxl.load_workbook(io.BytesIO(bundle.xlsx_bytes))
    joined = " ".join(wb.sheetnames).lower()
    for key in ("split", "cinema", "contrib", "absolute", "prisma"):
        assert key in joined, wb.sheetnames
    assert wb["PRISMA-NMA"].max_row == 33 and "PRISMA-NMA" in text


def test_figures_are_png_and_reasonably_sized() -> None:
    document, model, summary, result = basic()
    png = nma_report.network_png(model, summary, size=600)
    assert png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) < 400_000
    forest = nma_report.forest_png(result.random, result.measure, result.reference,
                                   sup.ranking_order(result.random), "95% CI", width=800)
    assert forest[:8] == b"\x89PNG\r\n\x1a\n"


def test_trace_png_from_chain_draws() -> None:
    import numpy as np

    rng = np.random.default_rng(0)
    draws = rng.normal(size=(3, 200, 2))
    tau = np.abs(rng.normal(size=(3, 200)))
    png = nma_report.trace_png(draws, ["d[A]", "d[B]"], tau, width=800)
    assert png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) > 1000
    assert nma_report.trace_png(draws, ["d[A]", "d[B]"], None, width=800)[:8] == b"\x89PNG\r\n\x1a\n"
    assert nma_report.trace_png(draws[:, :, :0], [], None) == b""


def test_transitivity_figure_is_drawn_and_reaches_the_report() -> None:
    document, model, summary, result = basic()
    document = {**document, "study_characteristics": [
        {"study": "Study 01", "mean_age": 62, "female_pct": 40, "year": 2015, "rob": "High"},
        {"study": "Study 02", "mean_age": 66, "female_pct": 52, "year": 2019, "rob": "Low"},
        {"study": "Study 03", "mean_age": 58, "female_pct": 47, "year": 2012},
    ]}
    points = sup.transitivity_points(summary, document["study_characteristics"], sup.characteristic_variables(document))
    png = nma_report.transitivity_png(points)
    assert png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) > 2000
    assert nma_report.transitivity_png(points.iloc[0:0]) == b""

    bundle = nma_report.build_report(nma_report.ReportInput(
        document=document, model=model, summary=summary, result=result, title="Transitivity"))
    assert bundle.figures["transitivity"][:8] == b"\x89PNG\r\n\x1a\n"
    text = _docx_text(bundle.docx_bytes)
    assert "Figure 1b" in text and "transitivity" in text
    # no characteristics -> no figure and no caption
    plain = nma_report.build_report(nma_report.ReportInput(
        document={**document, "study_characteristics": []}, model=model, summary=summary, result=result))
    assert "transitivity" not in plain.figures and "Figure 1b" not in _docx_text(plain.docx_bytes)
