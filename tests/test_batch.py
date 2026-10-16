import json

import pytest

from heor_tree_engine import batch


def _tree_doc(scenarios=None) -> dict:
    doc = {
        "root_id": "d",
        "parameters": [{"name": "p", "base": 0.2}, {"name": "c", "base": 100.0}],
        "nodes": [
            {"id": "d", "type": "decision", "branches": [
                {"label": "A", "child_id": "ca"}, {"label": "B", "child_id": "tb"}]},
            {"id": "ca", "type": "chance", "branches": [
                {"label": "event", "child_id": "t1", "probability": "p", "cost_increment": "c", "effect_increment": 0.0},
                {"label": "none", "child_id": "t2", "probability": "#", "cost_increment": 0.0, "effect_increment": 0.0}]},
            {"id": "t1", "type": "terminal", "cost": 0.0, "effect": 1.0, "year": 0.0, "branches": []},
            {"id": "t2", "type": "terminal", "cost": 0.0, "effect": 2.0, "year": 0.0, "branches": []},
            {"id": "tb", "type": "terminal", "cost": 10.0, "effect": 1.5, "year": 0.0, "branches": []},
        ],
    }
    if scenarios is not None:
        doc["sensitivity_analysis"] = {"willingness_to_pay": 1000.0, "scenarios": scenarios}
    return doc


def _markov_doc() -> dict:
    return {
        "id": "m", "name": "m", "cycles": 3, "cycles_per_year": 1.0, "half_cycle_correction": False,
        "parameters": [{"name": "pd", "base": 0.5}],
        "strategies": [{"id": "s1", "name": "S1", "states": [
            {"id": "well", "name": "Well", "cost": 100.0, "effect": 1.0, "initial_prob": 1.0},
            {"id": "dead", "name": "Dead", "cost": 0.0, "effect": 0.0, "initial_prob": 0.0, "is_absorbing": True}],
            "transitions": [
                {"from_state": "well", "to_state": "dead", "probability": "pd"},
                {"from_state": "well", "to_state": "well", "probability": "#"}]}],
    }


def test_detect_kind() -> None:
    assert batch.detect_kind(_tree_doc()) == "tree"
    assert batch.detect_kind(_markov_doc()) == "markov"
    assert batch.detect_kind({**_markov_doc(), "n_patients": 100}) == "microsim"
    assert batch.detect_kind({"studies": []}) is None
    assert batch.detect_kind({"foo": 1}) is None


def test_tree_base_case_and_saved_scenarios() -> None:
    doc = _tree_doc({"Base case": {}, "Cheap": {"c": 10.0}, "Likely": {"p": 0.9}})
    out = batch.run_document(doc, "t.decide")
    assert out.messages == [] and out.status == "OK"
    scenarios = [r["Scenario"] for r in out.rows]
    assert scenarios.count("Base case") == 2  # one row per strategy, base case only once
    assert set(scenarios) == {"Base case", "Cheap", "Likely"}
    base_a = next(r for r in out.rows if r["Scenario"] == "Base case" and r["Strategy"] == "A")
    assert base_a["Total cost"] == pytest.approx(0.2 * 100.0)
    assert base_a["Total effect"] == pytest.approx(0.2 * 1.0 + 0.8 * 2.0)
    cheap_a = next(r for r in out.rows if r["Scenario"] == "Cheap" and r["Strategy"] == "A")
    assert cheap_a["Total cost"] == pytest.approx(0.2 * 10.0)
    assert base_a["WTP"] == 1000.0


def test_markov_scenario_override_changes_result() -> None:
    doc = _markov_doc()
    doc["sensitivity_analysis"] = {"scenarios": {"Safer": {"pd": 0.1}}}
    out = batch.run_document(doc, "m.decide")
    assert out.messages == []
    base = next(r for r in out.rows if r["Scenario"] == "Base case")
    safer = next(r for r in out.rows if r["Scenario"] == "Safer")
    assert safer["Total cost"] > base["Total cost"]
    # cohort at start of cycles 0..2 with pd=0.5: 1 + 0.5 + 0.25 well-years x $100
    assert base["Total cost"] == pytest.approx(175.0)


def test_failing_scenario_does_not_stop_the_others() -> None:
    doc = _tree_doc({"Bad": {"does_not_exist": 1.0}, "Good": {"p": 0.5}})
    out = batch.run_document(doc, "t.decide")
    assert any("Bad" in m for m in out.messages)
    assert {r["Scenario"] for r in out.rows} == {"Base case", "Good"}
    assert out.status == "OK with errors"


def test_unsupported_and_unreadable_files(tmp_path) -> None:
    nma = tmp_path / "n.decide"
    nma.write_text(json.dumps({"studies": []}))
    bad = tmp_path / "bad.decide"
    bad.write_text("{not json")
    assert batch.run_file(str(nma)).status == "Failed"
    out = batch.run_file(str(bad))
    assert out.status == "Failed" and "Could not read" in out.messages[0]


def test_batch_over_folder_writes_excel(tmp_path) -> None:
    (tmp_path / "a.decide").write_text(json.dumps(_tree_doc({"Likely": {"p": 0.9}})))
    (tmp_path / "b.json").write_text(json.dumps(_markov_doc()))
    (tmp_path / "skip.txt").write_text("x")
    files = batch.collect_files([str(tmp_path)])
    assert [f.rsplit("/", 1)[-1] for f in files] == ["a.decide", "b.json"]
    outcomes = batch.run_batch(files, workers=2)
    assert [o.path for o in outcomes] == ["a.decide", "b.json"]
    out_file = tmp_path / "out.xlsx"
    batch.write_excel(outcomes, str(out_file))
    from openpyxl import load_workbook
    wb = load_workbook(out_file)
    assert wb.sheetnames == ["Results", "Files"]
    header = [c.value for c in wb["Results"][1]]
    assert header == list(batch.RESULT_COLUMNS)
    assert wb["Results"].max_row == 1 + sum(len(o.rows) for o in outcomes)
    assert [r[2].value for r in wb["Files"].iter_rows(min_row=2)] == ["OK", "OK"]
    assert batch.main([str(tmp_path / "a.decide"), "-o", str(tmp_path / "cli.xlsx"), "--workers", "1"]) == 0
    assert (tmp_path / "cli.xlsx").exists()
