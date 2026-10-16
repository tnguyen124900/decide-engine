"""Tests for microsim_report: report builds, self-checks pass on a sound
model, and a tampered result is caught."""

import copy
import io
import math
import sys
import zipfile
from pathlib import Path
from xml.dom import minidom

import pytest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import microsim_report as mr
from heor_tree_engine import resolve_markov, resolve_microsim, run_markov, run_microsim
from heor_tree_engine.selfcheck import SelfCheckFailure


def _doc(n=400):
    return {
        "id": "m", "name": "Toy", "cycles": 6, "half_cycle_correction": False, "n_patients": n, "random_seed": 3,
        "states": [
            {"id": "A", "name": "Well", "cost": 10.0, "effect": 1.0, "initial_prob": 1.0},
            {"id": "B", "name": "Ill", "cost": 50.0, "effect": 0.5, "initial_prob": 0.0},
            {"id": "D", "name": "Dead", "cost": 0.0, "effect": 0.0, "initial_prob": 0.0, "is_absorbing": True},
        ],
        "transitions": [
            {"from_state": "A", "to_state": "B", "probability": 0.2},
            {"from_state": "A", "to_state": "D", "probability": 0.05},
            {"from_state": "A", "to_state": "A", "probability": "#"},
            {"from_state": "B", "to_state": "A", "probability": 0.1},
            {"from_state": "B", "to_state": "D", "probability": 0.15},
            {"from_state": "B", "to_state": "B", "probability": "#"},
        ],
    }


def _input(n=400, cohort=True):
    doc = _doc(n)
    model = resolve_microsim(doc)
    res = run_microsim(model, {}, discount_cost_per_cycle=0.03, discount_effect_per_cycle=0.03)
    cohort_totals = None
    if cohort:
        mk = run_markov(resolve_markov(doc), {}, discount_cost_per_cycle=0.03, discount_effect_per_cycle=0.03)
        cohort_totals = (mk.total_cost, mk.total_effect)
    run = mr.MicrosimStrategyRun(
        label="Usual care", model=model, result=res, state_names={"A": "Well", "B": "Ill", "D": "Dead"},
        cohort_totals=cohort_totals, repro_totals=((1.5, 2.5), (1.5, 2.5)),
    )
    run2 = mr.MicrosimStrategyRun(
        label="Drug", model=model, result=res, state_names={"A": "Well", "B": "Ill", "D": "Dead"},
        cohort_totals=cohort_totals, repro_totals=((1.5, 2.5), (1.5, 2.5)),
    )
    return mr.MicrosimReportInput(
        model_name="Toy", parameters={}, runs=[run, run2], cost_symbol="$", effect_symbol="QALY",
        cost_decimals=2, effect_decimals=3, baseline_strategy="Usual care", maximize_effect=True,
        willingness_to_pay=50000.0, cost_discount_rate=0.03, effect_discount_rate=0.03,
    )


def test_report_builds_and_docx_is_wellformed():
    bundle = mr.build_microsim_report(_input())
    with zipfile.ZipFile(io.BytesIO(bundle.docx_bytes)) as zf:
        minidom.parseString(zf.read("word/document.xml"))
    assert bundle.xlsx_bytes[:2] == b"PK"
    assert bundle.filename_stem.startswith("decide-microsim-technical-report")


def test_selfchecks_pass_on_sound_model():
    checks = mr.microsim_selfchecks(_input(n=2000))
    failing = [c for c in checks if c.applicable and not c.passed]
    assert not failing, [(c.name, c.detail) for c in failing]
    assert sum(1 for c in checks if c.applicable) >= 6


def test_tampered_total_is_caught():
    inp = _input()
    inp.runs[0].result = copy.deepcopy(inp.runs[0].result)
    inp.runs[0].result.total_cost += 5.0
    with pytest.raises(SelfCheckFailure):
        mr.build_microsim_report(inp)


def test_cohort_disagreement_is_a_warning_not_a_block():
    inp = _input(n=2000)
    cost, effect = inp.runs[0].cohort_totals
    inp.runs[0].cohort_totals = (cost * 2, effect)
    checks = mr.microsim_selfchecks(inp)
    agree = [c for c in checks if "cohort" in c.name][0]
    assert agree.severity == "warning" and not agree.passed
    mr.build_microsim_report(inp)  # must still build


def test_icer_is_baseline_relative():
    inp = _input()
    rows = mr._icer_rows(inp)
    assert rows[0]["status"] == "-" and rows[0]["icer"] is None
    # identical results for both strategies -> no incremental effect -> dominated
    assert rows[1]["status"] == "Dominated"


def test_mcse_matches_definition():
    values = [1.0, 2.0, 3.0, 4.0]
    expected = math.sqrt(sum((v - 2.5) ** 2 for v in values) / 3 / 4)
    assert mr._mcse(values) == pytest.approx(expected)
    assert mr._mcse([1.0]) is None


def _docx_text(bundle):
    with zipfile.ZipFile(io.BytesIO(bundle.docx_bytes)) as zf:
        return zf.read("word/document.xml").decode("utf-8")


def test_report_has_cba_section_when_selected():
    inp = _input()
    inp.analysis_settings = {"type": "CBA", "cba_value_per_effect": 50000.0, "cba_extra_benefits": {"Drug": 100.0}}
    bundle = mr.build_microsim_report(inp)
    text = _docx_text(bundle)
    assert "Cost-benefit analysis" in text and "Net benefit" in text
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(bundle.xlsx_bytes))
    assert "CMA-CBA" in wb.sheetnames


def test_report_has_cma_section_and_default_has_none():
    inp = _input()
    inp.analysis_settings = {"type": "CMA", "cma_tolerance_pct": 1.0}
    text = _docx_text(mr.build_microsim_report(inp))
    assert "Cost-minimization analysis" in text
    inp.analysis_settings = {"type": "CEA"}
    bundle = mr.build_microsim_report(inp)
    assert "Cost-minimization" not in _docx_text(bundle) and "Cost-benefit" not in _docx_text(bundle)
    import openpyxl
    assert "CMA-CBA" not in openpyxl.load_workbook(io.BytesIO(bundle.xlsx_bytes)).sheetnames
