"""Shared, Streamlit-free report-writing building blocks: a raw-OOXML
``.docx`` writer (no ``python-docx`` dependency) and a multi-sheet ``.xlsx``
exporter built on ``openpyxl``/``pandas``.

Extracted from ``nma_report.py`` so the Tree and Markov Technical Documentation
Reports can reuse the exact same Word/Excel plumbing the NMA report already
relies on, without importing the NMA module (which pulls in NMA-specific
dependencies that have nothing to do with Tree/Markov reports).

This is a pure extraction: the two pieces below (``Docx`` and
``write_xlsx``) are byte-for-byte the same implementation that used to live
in ``nma_report.py`` as ``_Docx``/``write_xlsx``. ``nma_report.py`` now
imports them from here instead of defining them itself, so NMA's own report
output is unchanged.

The one behavioural difference from the original: ``write_xlsx`` used to
read a module-level ``SHEET_ORDER`` constant that only made sense for NMA's
own report tables (Effects/League table/SOF/...). It now takes
``sheet_order`` as an explicit parameter, so a caller building a Tree or
Markov report passes its own (key, sheet name) list instead.
"""

from __future__ import annotations

import io
import zipfile
from typing import Sequence
from xml.sax.saxutils import escape

import pandas as pd

# Dark-blue used for table header rows in both the .docx and the .xlsx --
# keeps the two documents visually consistent with each other.
HEADER_BG = "1E3A8A"


def write_xlsx(
    tables: dict[str, pd.DataFrame],
    title: str,
    methods: list[str],
    sheet_order: Sequence[tuple[str, str]],
) -> bytes:
    """One .xlsx with a cover sheet ("Report": title + methods paragraphs)
    followed by one sheet per entry in ``sheet_order`` -- ``(key, sheet
    name)`` pairs, in the order they should appear. A key missing from
    ``tables``, or present but empty, is skipped (no blank sheet). Every
    data sheet gets a bold white-on-navy header row, freeze pane under the
    header, and column widths sized to fit their content (capped at 60
    characters so one long value can't blow out the whole sheet)."""

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        cover = pd.DataFrame({"": [title, "", *methods]})
        cover.to_excel(writer, sheet_name="Report", index=False, header=False)
        writer.sheets["Report"].column_dimensions["A"].width = 120
        for row in writer.sheets["Report"].iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        writer.sheets["Report"]["A1"].font = Font(bold=True, size=14)
        for key, name in sheet_order:
            frame = tables.get(key)
            if frame is None or frame.empty:
                continue
            frame.to_excel(writer, sheet_name=name[:31], index=False)
            sheet = writer.sheets[name[:31]]
            fill = PatternFill("solid", fgColor=HEADER_BG)
            for cell in sheet[1]:
                cell.fill = fill
                cell.font = Font(bold=True, color="FFFFFF")
            for i, column in enumerate(frame.columns, start=1):
                longest = max([len(str(column))] + [len(str(v)) for v in frame[column].tolist()[:200]])
                sheet.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
            sheet.freeze_panes = "A2"
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Word (.docx written directly as OOXML)
# --------------------------------------------------------------------------


class Docx:
    """Minimal .docx builder: headings, paragraphs, captions, bullets,
    page breaks, a styled table, and an embedded PNG image -- everything a
    results-section report needs, written directly as OOXML so there is no
    ``python-docx`` dependency."""

    def __init__(self) -> None:
        self.body: list[str] = []
        self.images: list[tuple[str, bytes]] = []  # (name, png)

    # -- text ----------------------------------------------------------
    @staticmethod
    def _run(text: str, bold: bool = False, italic: bool = False, size: int | None = None, color: str | None = None) -> str:
        props = ""
        if bold:
            props += "<w:b/>"
        if italic:
            props += "<w:i/>"
        if color:
            props += f'<w:color w:val="{color}"/>'
        if size:
            props += f'<w:sz w:val="{size * 2}"/>'
        rpr = f"<w:rPr>{props}</w:rPr>" if props else ""
        return f'<w:r>{rpr}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'

    def heading(self, text: str, level: int = 1) -> None:
        self.body.append(f'<w:p><w:pPr><w:pStyle w:val="Heading{level}"/></w:pPr>{self._run(text)}</w:p>')

    def paragraph(self, text: str, italic: bool = False, size: int | None = None, align: str | None = None) -> None:
        ppr = f'<w:pPr><w:jc w:val="{align}"/></w:pPr>' if align else ""
        self.body.append(f"<w:p>{ppr}{self._run(text, italic=italic, size=size)}</w:p>")

    def caption(self, text: str) -> None:
        self.body.append(f'<w:p><w:pPr><w:pStyle w:val="Caption"/></w:pPr>{self._run(text, italic=True, size=9)}</w:p>')

    def bullet(self, text: str) -> None:
        self.body.append(f'<w:p><w:pPr><w:ind w:left="360" w:hanging="360"/></w:pPr>{self._run("• " + text)}</w:p>')

    def page_break(self) -> None:
        self.body.append('<w:p><w:r><w:br w:type="page"/></w:r></w:p>')

    # -- table ---------------------------------------------------------
    def table(self, frame: pd.DataFrame, max_rows: int = 400, font_size: int = 8,
              widths: list[float] | None = None, highlight_rows: "set[int] | None" = None,
              highlight_fill: str = "FBDADA") -> None:
        """``highlight_rows``: 0-based row indices (into ``frame``, before
        the ``max_rows`` truncation) to mark with ``highlight_fill``
        instead of the normal alternating stripe -- e.g. a parameter
        missing its Source citation, so a reviewer sees at a glance which
        rows still need one."""

        highlight_rows = highlight_rows or set()
        columns = [str(c) for c in frame.columns]
        if widths is None:
            widths = []
            for column in columns:
                longest = max([len(column)] + [len(str(v)) for v in frame[column].tolist()[:max_rows]])
                widths.append(min(40, max(8, longest)))
        total = sum(widths) or 1
        twips = [int(9000 * w / total) for w in widths]
        cells_head = "".join(
            f'<w:tc><w:tcPr><w:tcW w:w="{tw}" w:type="dxa"/><w:shd w:val="clear" w:color="auto" w:fill="{HEADER_BG}"/></w:tcPr>'
            f'<w:p><w:pPr><w:spacing w:before="20" w:after="20"/></w:pPr>{self._run(column, bold=True, size=font_size, color="FFFFFF")}</w:p></w:tc>'
            for column, tw in zip(columns, twips))
        rows_xml = [f"<w:tr><w:trPr><w:tblHeader/></w:trPr>{cells_head}</w:tr>"]
        for r, (_, row) in enumerate(frame.head(max_rows).iterrows()):
            if r in highlight_rows:
                fill = f' w:fill="{highlight_fill}"'
            else:
                fill = ' w:fill="F5F8FE"' if r % 2 else ' w:fill="FFFFFF"'
            cells = "".join(
                f'<w:tc><w:tcPr><w:tcW w:w="{tw}" w:type="dxa"/><w:shd w:val="clear" w:color="auto"{fill}/></w:tcPr>'
                f'<w:p><w:pPr><w:spacing w:before="20" w:after="20"/></w:pPr>{self._run("" if pd.isna(v) else str(v), size=font_size)}</w:p></w:tc>'
                for v, tw in zip(row.tolist(), twips))
            rows_xml.append(f"<w:tr>{cells}</w:tr>")
        borders = "".join(f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="DBE2EE"/>'
                          for side in ("top", "left", "bottom", "right", "insideH", "insideV"))
        grid = "".join(f'<w:gridCol w:w="{tw}"/>' for tw in twips)
        self.body.append(
            f'<w:tbl><w:tblPr><w:tblW w:w="9000" w:type="dxa"/><w:tblBorders>{borders}</w:tblBorders>'
            f'<w:tblLayout w:type="fixed"/></w:tblPr><w:tblGrid>{grid}</w:tblGrid>{"".join(rows_xml)}</w:tbl>')
        self.body.append("<w:p/>")

    # -- image ---------------------------------------------------------
    def image(self, png: bytes, width_cm: float = 16.0) -> None:
        from PIL import Image

        with Image.open(io.BytesIO(png)) as im:
            w, h = im.size
        index = len(self.images) + 1
        name = f"image{index}.png"
        self.images.append((name, png))
        cx = int(width_cm * 360000)
        cy = int(cx * h / w)
        rid = f"rIdImg{index}"
        self.body.append(
            f'<w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:drawing>'
            f'<wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="{cx}" cy="{cy}"/>'
            f'<wp:docPr id="{index}" name="{name}"/>'
            f'<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f'<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f'<pic:nvPicPr><pic:cNvPr id="{index}" name="{name}"/><pic:cNvPicPr/></pic:nvPicPr>'
            f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
            f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic>'
            f"</wp:inline></w:drawing></w:r></w:p>")

    # -- package -------------------------------------------------------
    def to_bytes(self) -> bytes:
        ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
              'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
              'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"')
        document = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {ns}><w:body>'
                    + "".join(self.body)
                    + '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1300" w:right="1300" w:bottom="1300" '
                      'w:left="1300" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr></w:body></w:document>')
        styles = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                  '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri"/>'
                  '<w:sz w:val="21"/></w:rPr></w:rPrDefault><w:pPrDefault><w:pPr><w:spacing w:after="120" w:line="276" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>'
                  '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
                  '<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/>'
                  '<w:pPr><w:spacing w:after="200"/></w:pPr><w:rPr><w:b/><w:color w:val="1E3A8A"/><w:sz w:val="40"/></w:rPr></w:style>'
                  '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/>'
                  '<w:pPr><w:keepNext/><w:spacing w:before="320" w:after="120"/><w:outlineLvl w:val="0"/></w:pPr>'
                  '<w:rPr><w:b/><w:color w:val="1E3A8A"/><w:sz w:val="30"/></w:rPr></w:style>'
                  '<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/>'
                  '<w:pPr><w:keepNext/><w:spacing w:before="240" w:after="80"/><w:outlineLvl w:val="1"/></w:pPr>'
                  '<w:rPr><w:b/><w:color w:val="1E3A8A"/><w:sz w:val="25"/></w:rPr></w:style>'
                  '<w:style w:type="paragraph" w:styleId="Caption"><w:name w:val="caption"/><w:basedOn w:val="Normal"/>'
                  '<w:pPr><w:spacing w:before="60" w:after="200"/></w:pPr><w:rPr><w:i/><w:color w:val="475569"/><w:sz w:val="18"/></w:rPr></w:style>'
                  '</w:styles>')
        rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                '</Relationships>')
        doc_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    '<Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                    + "".join(f'<Relationship Id="rIdImg{i + 1}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/{name}"/>'
                              for i, (name, _) in enumerate(self.images))
                    + "</Relationships>")
        content_types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                         '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                         '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                         '<Default Extension="xml" ContentType="application/xml"/>'
                         '<Default Extension="png" ContentType="image/png"/>'
                         '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                         '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
                         '</Types>')
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", content_types)
            zf.writestr("_rels/.rels", rels)
            zf.writestr("word/document.xml", document)
            zf.writestr("word/styles.xml", styles)
            zf.writestr("word/_rels/document.xml.rels", doc_rels)
            for name, png in self.images:
                zf.writestr(f"word/media/{name}", png)
        return buffer.getvalue()


# Validation self-checks (shared plumbing lives in the package so engines can build checks without importing this
# root-level module); re-exported here for the report writers.
from heor_tree_engine.selfcheck import (  # noqa: E402
    DEFAULT_ABS_TOL, DEFAULT_REL_TOL, SelfCheckFailure, SelfCheckResult, raise_on_failure, values_close,
)


def selfchecks_frame(checks: "Sequence[SelfCheckResult]") -> pd.DataFrame:
    return pd.DataFrame([
        {"Check": c.name,
         "Result": ("Pass" if c.passed else ("FAIL" if c.severity == "error" else "Review")) if c.applicable else "N/A",
         "Detail": c.detail}
        for c in checks
    ])


def write_selfchecks_section(doc: "Docx", checks: "Sequence[SelfCheckResult]", heading: str) -> None:
    """Heading + scope paragraph + Check / Result / Detail table."""

    doc.heading(heading, level=1)
    doc.paragraph(
        "The checks below were run automatically while this report was built, on this exact model. They test the "
        "internal validity of the programmed model (consistency and extreme-value tests, in the sense of the "
        "AdViSHE tool's 'programmed model' items). They do not replace conceptual validation, validation of the "
        "input data, comparison with independent data or models, or expert face-validity review, which cannot be "
        "automated. A failing check would have stopped this report from being produced; 'Review' marks a model-quality finding that does not stop it."
    )
    doc.table(selfchecks_frame(checks), widths=[26, 8, 60])
