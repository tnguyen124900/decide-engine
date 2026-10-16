"""Excel export of a Decision Tree model with a live, auditable formula
trail -- built for HTA submissions, where a reviewer
must be able to open the file, trace every number back to a parameter or
a listed branch value, and re-run the analysis by editing a cell.

Two-tier transparency: the ``Parameters`` sheet holds *resolved
values* -- a parameter may come from an arbitrary DECIDE expression
(distributions, tables, DistMean, ...) with no 1:1 Excel equivalent for
every case, so its Base is a plain number, with the parameter's own
Distribution/Dist args/Source alongside for documentation, not a live
formula. Everything the *engine* does on top of those parameters --
summing a path's branch increments, weighting paths by probability,
rolling paths up into a strategy, computing an ICER between two
strategies -- IS a live Excel formula, because that part is exactly the
engine's own arithmetic and every step of it is expressible as a
spreadsheet formula.

Reuses ``tree_paths.enumerate_paths()``'s ``PathRow``/``PathStep`` -- the
exact same per-path numbers already used by the app's own "per-path
results" panel and asserted (tests/test_tree_paths.py) to reconcile with
``Engine.run()``'s expected cost/effect -- so this export cannot silently
drift from what the app shows on screen.

Discounting note: per ``tree_paths.py``'s own docstring, a branch's
cost/effect increment is NEVER discounted -- only a leaf's payoff is (by
the leaf's own ``year``). So every non-leaf step has identical
undiscounted/discounted values, and this module only asks the leaf step
for both variants (see ``_leaf_step``); the "Cumulative" formulas for
both undiscounted and discounted totals sum the *same* non-leaf step
columns plus whichever leaf figure applies.

Extra outcomes (Outcomes-table rows beyond the primary Cost/Effect pair)
are included automatically whenever the model has them (mirrors the
engine's own ``extra_outcome_ids`` convention -- nothing to configure),
but as **resolved per-path totals** rather than a live per-step
breakdown: doing that fully live would double the sheet's width for a
case that is, in practice, rare. Cells carrying such a resolved (not
live) value get a comment saying so, following the rule that a
non-live cell must say so rather than look like a formula result.
"""

from __future__ import annotations

import hashlib
import io
import re
from collections.abc import Mapping, Sequence
from datetime import datetime, timezone
from typing import Any

from .models import Param, TerminalNode
from .tree_paths import PathRow

_HEADER_FILL = "1E3A8A"
_HEADER_FONT = "FFFFFF"
_RESOLVED_FILL = "FFF3CD"  # pale amber: "this is a resolved value, not a live formula"
_MISSING_SOURCE_FILL = "FDE2E2"  # pale red: parameter has no Source citation


def _excel_safe_defined_name(raw_name: str, used: set[str]) -> str:
    """A valid, unique Excel defined name for parameter ``raw_name``.

    Excel defined names must start with a letter/underscore, contain only
    letters/digits/underscores/periods/backslashes, and must not look like
    a cell reference (e.g. "A1"). Prefixing with "p_" sidesteps both the
    cell-reference collision and names starting with a digit; anything
    else invalid is replaced with "_", and a numeric suffix is added if
    two different parameter names collide after sanitising.
    """

    cleaned = re.sub(r"[^0-9A-Za-z_]", "_", raw_name)
    candidate = f"p_{cleaned}" if cleaned else "p_param"
    if not re.match(r"^[A-Za-z_]", candidate):
        candidate = f"p_{candidate}"
    base = candidate
    suffix = 2
    while candidate in used:
        candidate = f"{base}_{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def _excel_quote(text: str) -> str:
    """A string literal for use inside an Excel formula (doubles embedded
    double-quotes, per Excel's own escaping rule)."""

    return '"' + str(text).replace('"', '""') + '"'


def _config_hash(parameters: Mapping[str, Param], paths: Sequence[PathRow]) -> str:
    """A short, stable fingerprint of the exported model configuration, so
    a company can later prove a given Excel file matches a given DECIDE
    model state. Not a security hash -- just a change-detector."""

    parts: list[str] = []
    for name in sorted(parameters):
        p = parameters[name]
        parts.append(f"{name}={p.base}|{p.low}|{p.high}|{p.dist}|{p.dist_args}|{p.source}")
    for row in paths:
        parts.append(f"path={row.strategy}|{row.path_id}|{row.p}|{row.discounted}")
    digest = hashlib.sha256("\n".join(parts).encode("utf-8")).hexdigest()
    return digest[:12]


def _style_header_row(ws: Any, row: int, n_cols: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor=_HEADER_FILL)
    font = Font(bold=True, color=_HEADER_FONT)
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def _autosize(ws: Any, n_cols: int, sample_rows: int = 200) -> None:
    from openpyxl.utils import get_column_letter

    for col in range(1, n_cols + 1):
        letter = get_column_letter(col)
        longest = 8
        for row in ws.iter_rows(min_col=col, max_col=col, max_row=sample_rows):
            for cell in row:
                if cell.value is not None:
                    longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(46, max(10, longest + 2))


def _write_parameters_sheet(
    wb: Any, parameters: Mapping[str, Param]
) -> dict[str, dict[str, Any]]:
    """Write the ``Parameters`` sheet; return, per parameter name, its
    Excel defined name and the 1-based row it lives on (for cross-sheet
    formulas and for the missing-Source pre-flight check)."""

    ws = wb.active
    ws.title = "Parameters"
    headers = ["Name", "Base", "Low", "High", "Distribution", "Dist args", "Source"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    used_names: set[str] = set()
    info: dict[str, dict[str, Any]] = {}
    row_index = 2
    for name in sorted(parameters):
        param = parameters[name]
        defined_name = _excel_safe_defined_name(name, used_names)
        ws.cell(row=row_index, column=1, value=name)
        base_cell = ws.cell(row=row_index, column=2, value=float(param.base))
        ws.cell(row=row_index, column=3, value=param.low)
        ws.cell(row=row_index, column=4, value=param.high)
        ws.cell(row=row_index, column=5, value=param.dist)
        ws.cell(
            row=row_index,
            column=6,
            value=(", ".join(f"{k}={v}" for k, v in param.dist_args.items()) if param.dist_args else None),
        )
        source_cell = ws.cell(row=row_index, column=7, value=param.source)
        if not param.source:
            from openpyxl.styles import PatternFill

            source_cell.fill = PatternFill("solid", fgColor=_MISSING_SOURCE_FILL)

        from openpyxl.workbook.defined_name import DefinedName

        wb.defined_names[defined_name] = DefinedName(
            defined_name, attr_text=f"Parameters!${base_cell.column_letter}${row_index}"
        )
        info[name] = {"defined_name": defined_name, "row": row_index}
        row_index += 1

    _autosize(ws, len(headers), sample_rows=row_index)
    return info


def _leaf_step(row: PathRow):
    return row.steps[-1]


def _non_leaf_steps(row: PathRow):
    return row.steps[:-1]


def _cell_live_or_resolved(
    ws: Any, row: int, col: int, raw_ref: Any, resolved_value: float,
    param_defined_names: Mapping[str, str], resolved_fill: Any,
) -> Any:
    """Write one cell as a live formula when ``raw_ref`` is exactly a known
    parameter's name (``_live_ref``), otherwise as the already-resolved
    number, filled/commented as "resolved, not live" -- the one repeated
    decision this sheet makes over and over, factored out so every call
    site (probability, cost, effect, leaf cost/effect) applies it the same
    way."""

    live = _live_ref(raw_ref, param_defined_names)
    cell = ws.cell(row=row, column=col, value=live if live is not None else float(resolved_value))
    # Only an *expression* (a string DECIDE had to evaluate) that isn't a
    # bare parameter reference needs the "resolved, not live" disclosure --
    # a plain literal number in the model (e.g. a branch probability typed
    # in directly as 0.6) has nothing hidden behind it in the first place.
    if live is None and isinstance(raw_ref, str):
        cell.fill = resolved_fill
        cell.comment = _resolved_comment()
    return cell


def _write_tree_paths_sheet(
    wb: Any,
    tree: Any,
    paths: Sequence[PathRow],
    param_defined_names: Mapping[str, str],
    cost_discount_rate: float,
    effect_discount_rate: float,
    extra_outcome_ids: Sequence[str],
    extra_labels: Mapping[str, str],
    psm_links: Mapping[str, Mapping[str, str]] | None = None,
) -> dict[str, Any]:
    """Write the ``Tree Paths`` sheet: one row per root-to-leaf path, with
    the joint probability and every cost/effect contribution built as a
    live formula from the step-by-step columns to its left -- see the
    module docstring for what stays "resolved value" vs "live formula".
    A step's own probability/cost/effect is ALSO live whenever the model
    set it to a bare parameter reference (e.g. a branch cost of exactly
    "cost_drug", or a chance branch probability of exactly "p_resp") --
    see ``_live_ref`` -- which covers the common case in real models,
    not just the totals built on top of it.

    Returns column/row bookkeeping the ``Strategy Summary`` sheet needs:
    the header row, first/last data row, and the column letters for
    Strategy / Expected cost contribution / Expected effect contribution
    / each extra outcome's expected contribution.
    """

    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("Tree Paths")
    max_non_leaf = max((len(_non_leaf_steps(row)) for row in paths), default=0)

    headers = ["Strategy", "Path"]
    step_col_start = len(headers) + 1
    for i in range(1, max_non_leaf + 1):
        headers += [f"Step {i} probability", f"Step {i} cost", f"Step {i} effect"]
    leaf_cols_start = len(headers) + 1
    headers += [
        "Leaf cost (undiscounted)", "Leaf cost (discounted)",
        "Leaf effect (undiscounted)", "Leaf effect (discounted)",
    ]
    joint_prob_col = len(headers) + 1
    headers += ["Joint probability"]
    cum_cols_start = len(headers) + 1
    headers += [
        "Total cost (undiscounted)", "Total cost (discounted)",
        "Total effect (undiscounted)", "Total effect (discounted)",
    ]
    expected_cols_start = len(headers) + 1
    headers += ["Expected cost contribution", "Expected effect contribution"]
    extra_cols_start = len(headers) + 1
    for oid in extra_outcome_ids:
        label = extra_labels.get(oid, oid)
        headers += [f"{label} total (discounted)", f"Expected {label} contribution"]

    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    resolved_fill = PatternFill("solid", fgColor=_RESOLVED_FILL)
    first_data_row = 2
    for offset, row in enumerate(paths):
        r = first_data_row + offset
        non_leaf = _non_leaf_steps(row)
        leaf = _leaf_step(row)
        branches, leaf_node = _resolve_path_branches(tree, row.labels)

        ws.cell(row=r, column=1, value=row.strategy)
        ws.cell(row=r, column=2, value=row.path_label)

        prob_cols: list[str] = []
        cost_cols: list[str] = []
        effect_cols: list[str] = []
        col = step_col_start
        for step, branch in zip(non_leaf, branches):
            prob_cell = _cell_live_or_resolved(
                ws, r, col, branch.probability if branch.probability is not None else None,
                step.p_step if step.p_step is not None else 0.0, param_defined_names, resolved_fill,
            ) if step.p_step is not None else ws.cell(row=r, column=col, value=None)
            cost_cell = _cell_live_or_resolved(
                ws, r, col + 1, branch.cost_increment, step.undiscounted.get("cost", 0.0),
                param_defined_names, resolved_fill,
            )
            effect_cell = _cell_live_or_resolved(
                ws, r, col + 2, branch.effect_increment, step.undiscounted.get("effect", 0.0),
                param_defined_names, resolved_fill,
            )
            if step.p_step is not None:
                prob_cols.append(prob_cell.coordinate)
            cost_cols.append(cost_cell.coordinate)
            effect_cols.append(effect_cell.coordinate)
            col += 3
        # Pad unused step slots (shorter paths) with blank, unfilled cells.
        for _ in range(max_non_leaf - len(non_leaf)):
            col += 3

        psm_link = (psm_links or {}).get(row.leaf_id) if row.leaf_type == "psm" else None
        is_terminal_leaf = isinstance(leaf_node, TerminalNode)
        leaf_cost_raw = leaf_node.cost if is_terminal_leaf else None
        leaf_effect_raw = leaf_node.effect if is_terminal_leaf else None
        leaf_year_raw = leaf_node.year if is_terminal_leaf else None
        year_is_literal = isinstance(leaf_year_raw, (int, float)) and not isinstance(leaf_year_raw, bool)

        if psm_link is not None:
            # A partitioned-survival leaf: its payoff is the node's own live sheet total, not a resolved number.
            leaf_cost_u = ws.cell(row=r, column=leaf_cols_start, value="=" + psm_link["cost_undiscounted"])
            leaf_cost_d = ws.cell(row=r, column=leaf_cols_start + 1, value="=" + psm_link["cost_discounted"])
            leaf_effect_u = ws.cell(row=r, column=leaf_cols_start + 2, value="=" + psm_link["effect_undiscounted"])
            leaf_effect_d = ws.cell(row=r, column=leaf_cols_start + 3, value="=" + psm_link["effect_discounted"])
        else:
            leaf_cost_u = _cell_live_or_resolved(
                ws, r, leaf_cols_start, leaf_cost_raw, leaf.undiscounted.get("cost", 0.0),
                param_defined_names, resolved_fill,
            )
            leaf_effect_u = _cell_live_or_resolved(
                ws, r, leaf_cols_start + 2, leaf_effect_raw, leaf.undiscounted.get("effect", 0.0),
                param_defined_names, resolved_fill,
            )
        # The discounted leaf figure is live only when BOTH the leaf's own
        # cost/effect is a bare parameter reference AND its discount year
        # is a plain literal (not itself an expression) -- otherwise the
        # discount math can't be expressed purely from cells already on
        # this sheet, so it falls back to the resolved number.
        if psm_link is None:
            cost_live = _live_ref(leaf_cost_raw, param_defined_names)
            if cost_live is not None and year_is_literal:
                leaf_cost_d = ws.cell(
                    row=r, column=leaf_cols_start + 1,
                    value=f"={leaf_cost_u.coordinate}/(1+{cost_discount_rate})^{float(leaf_year_raw)}",
                )
            else:
                leaf_cost_d = ws.cell(row=r, column=leaf_cols_start + 1, value=float(leaf.discounted.get("cost", 0.0)))
                leaf_cost_d.fill = resolved_fill
                leaf_cost_d.comment = _resolved_comment()
            effect_live = _live_ref(leaf_effect_raw, param_defined_names)
            if effect_live is not None and year_is_literal:
                leaf_effect_d = ws.cell(
                    row=r, column=leaf_cols_start + 3,
                    value=f"={leaf_effect_u.coordinate}/(1+{effect_discount_rate})^{float(leaf_year_raw)}",
                )
            else:
                leaf_effect_d = ws.cell(row=r, column=leaf_cols_start + 3, value=float(leaf.discounted.get("effect", 0.0)))
                leaf_effect_d.fill = resolved_fill
                leaf_effect_d.comment = _resolved_comment()

        joint_cell = ws.cell(row=r, column=joint_prob_col)
        joint_cell.value = ("=" + "*".join(prob_cols)) if prob_cols else 1.0

        cost_sum = "+".join(cost_cols) if cost_cols else "0"
        effect_sum = "+".join(effect_cols) if effect_cols else "0"
        total_cost_u = ws.cell(row=r, column=cum_cols_start, value=f"=({cost_sum})+{leaf_cost_u.coordinate}")
        total_cost_d = ws.cell(row=r, column=cum_cols_start + 1, value=f"=({cost_sum})+{leaf_cost_d.coordinate}")
        total_effect_u = ws.cell(row=r, column=cum_cols_start + 2, value=f"=({effect_sum})+{leaf_effect_u.coordinate}")
        total_effect_d = ws.cell(row=r, column=cum_cols_start + 3, value=f"=({effect_sum})+{leaf_effect_d.coordinate}")

        ws.cell(row=r, column=expected_cols_start, value=f"={joint_cell.coordinate}*{total_cost_d.coordinate}")
        ws.cell(row=r, column=expected_cols_start + 1, value=f"={joint_cell.coordinate}*{total_effect_d.coordinate}")

        col = extra_cols_start
        for oid in extra_outcome_ids:
            total_cell = ws.cell(row=r, column=col, value=float(row.discounted.get(oid, 0.0)))
            total_cell.fill = resolved_fill
            total_cell.comment = _resolved_comment()
            ws.cell(row=r, column=col + 1, value=f"={joint_cell.coordinate}*{total_cell.coordinate}")
            col += 2

    last_data_row = first_data_row + len(paths) - 1 if paths else first_data_row - 1
    _autosize(ws, len(headers))

    def col_letter(index: int) -> str:
        return get_column_letter(index)

    bookkeeping = {
        "sheet_name": "Tree Paths",
        "first_data_row": first_data_row,
        "last_data_row": last_data_row,
        "strategy_col": col_letter(1),
        "expected_cost_col": col_letter(expected_cols_start),
        "expected_effect_col": col_letter(expected_cols_start + 1),
        "extra_expected_cols": {
            oid: col_letter(extra_cols_start + 2 * i + 1) for i, oid in enumerate(extra_outcome_ids)
        },
    }
    return bookkeeping


def _resolved_comment():
    from openpyxl.comments import Comment

    return Comment(
        "Resolved value, not a live formula: DECIDE evaluated this from the model's own "
        "expression/distribution (see the Parameters sheet for named inputs).",
        "DECIDE",
    )


def _live_ref(raw_value: Any, param_defined_names: Mapping[str, str]) -> str | None:
    """If ``raw_value`` (a NumericRef: a float, or a string that is either
    a bare parameter name or a full expression) is EXACTLY one known
    parameter's name, return the live Excel formula referencing that
    parameter's defined name (e.g. ``"=cost_drug"``). Otherwise (a literal
    number, or an expression more complex than a bare name -- arithmetic,
    a distribution function, a table lookup...) return None, meaning the
    caller must fall back to the already-resolved numeric value. Bare
    parameter references are extremely common in real models (a branch's
    cost is very often simply "cost_drug"), so this one case alone makes
    most of a typical tree's Excel export genuinely live end-to-end; a
    fuller expression-to-Excel transpiler is out of scope (see the module
    docstring)."""

    if not isinstance(raw_value, str):
        return None
    name = raw_value.strip()
    defined_name = param_defined_names.get(name)
    return f"={defined_name}" if defined_name else None


def _resolve_path_branches(tree: Any, labels: Sequence[str]) -> tuple[list[Any], Any]:
    """Walk ``tree`` following ``labels`` (one label per Decision/Chance
    branch taken, root to the node just before the leaf -- exactly
    ``PathRow.labels``) and return the ``Branch`` objects traversed plus
    the leaf ``Node`` reached. Matches tree_paths.py's own walk, just
    replaying it by label instead of enumerating every branch, since here
    we already know which path we want."""

    branches: list[Any] = []
    node: Any = tree
    for label in labels:
        branch = next(b for b in node.branches if b.label == label)
        branches.append(branch)
        node = branch.child
    return branches, node


def _write_strategy_summary_sheet(
    wb: Any, paths_info: Mapping[str, Any], strategy_order: Sequence[str], extra_outcome_ids: Sequence[str],
    extra_labels: Mapping[str, str],
) -> dict[str, Any]:
    ws = wb.create_sheet("Strategy Summary")
    headers = ["Strategy", "Expected cost", "Expected effect"]
    for oid in extra_outcome_ids:
        headers.append(f"Expected {extra_labels.get(oid, oid)}")
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    sheet = paths_info["sheet_name"]
    r0, r1 = paths_info["first_data_row"], paths_info["last_data_row"]
    strategy_rng = f"'{sheet}'!${paths_info['strategy_col']}${r0}:${paths_info['strategy_col']}${r1}"

    rows_by_strategy: dict[str, int] = {}
    for offset, strategy in enumerate(strategy_order):
        r = 2 + offset
        ws.cell(row=r, column=1, value=strategy)
        crit = _excel_quote(strategy)
        cost_rng = f"'{sheet}'!${paths_info['expected_cost_col']}${r0}:${paths_info['expected_cost_col']}${r1}"
        effect_rng = f"'{sheet}'!${paths_info['expected_effect_col']}${r0}:${paths_info['expected_effect_col']}${r1}"
        ws.cell(row=r, column=2, value=f"=SUMIF({strategy_rng},{crit},{cost_rng})")
        ws.cell(row=r, column=3, value=f"=SUMIF({strategy_rng},{crit},{effect_rng})")
        col = 4
        for oid in extra_outcome_ids:
            extra_col = paths_info["extra_expected_cols"][oid]
            extra_rng = f"'{sheet}'!${extra_col}${r0}:${extra_col}${r1}"
            ws.cell(row=r, column=col, value=f"=SUMIF({strategy_rng},{crit},{extra_rng})")
            col += 1
        rows_by_strategy[strategy] = r

    _autosize(ws, len(headers))
    return {"sheet_name": "Strategy Summary", "rows_by_strategy": rows_by_strategy, "cost_col": "B", "effect_col": "C"}


def _write_icer_sheet(
    wb: Any,
    summary_info: Mapping[str, Any],
    strategy_order: Sequence[str],
    baseline_strategy: str | None,
    maximize_effect: bool,
    willingness_to_pay: float,
) -> None:
    """Mirrors the app's own CEA convention exactly (see
    ``_render_tree_analysis_section`` in app.py): every strategy compared
    against ONE baseline (not a sorted efficient-frontier with extended
    dominance), so this sheet's numbers reconcile 1:1 with what the app
    shows on screen for the same model."""

    ws = wb.create_sheet("ICER")
    headers = ["Strategy", "Expected cost", "Expected effect", "Incremental cost", "Incremental effect", "ICER", "Status"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    if not strategy_order:
        return
    baseline = baseline_strategy if baseline_strategy in strategy_order else strategy_order[0]
    sheet = summary_info["sheet_name"]
    rows_by_strategy = summary_info["rows_by_strategy"]
    base_row = rows_by_strategy[baseline]
    base_cost_ref = f"'{sheet}'!${summary_info['cost_col']}${base_row}"
    base_effect_ref = f"'{sheet}'!${summary_info['effect_col']}${base_row}"
    direction = 1 if maximize_effect else -1

    for offset, strategy in enumerate(strategy_order):
        r = 2 + offset
        src_row = rows_by_strategy[strategy]
        cost_ref = f"'{sheet}'!${summary_info['cost_col']}${src_row}"
        effect_ref = f"'{sheet}'!${summary_info['effect_col']}${src_row}"
        ws.cell(row=r, column=1, value=strategy)
        ws.cell(row=r, column=2, value=f"={cost_ref}")
        ws.cell(row=r, column=3, value=f"={effect_ref}")
        inc_cost_cell = ws.cell(row=r, column=4, value=f"={cost_ref}-{base_cost_ref}")
        inc_effect_cell = ws.cell(row=r, column=5, value=f"={effect_ref}-{base_effect_ref}")
        if strategy == baseline:
            ws.cell(row=r, column=6, value=None)
            ws.cell(row=r, column=7, value="-")
        else:
            ws.cell(
                row=r, column=6,
                value=f"=IF(OR({inc_effect_cell.coordinate}=0,{direction}*{inc_effect_cell.coordinate}<=0),"
                      f"NA(),{inc_cost_cell.coordinate}/{inc_effect_cell.coordinate})",
            )
            ws.cell(
                row=r, column=7,
                value=f'=IF(OR({inc_effect_cell.coordinate}=0,{direction}*{inc_effect_cell.coordinate}<=0),'
                      f'"Dominated",IF(F{r}<={willingness_to_pay},"Within WTP","Above WTP"))',
            )
    _autosize(ws, len(headers))


def _write_psm_sheets(
    wb: Any,
    tree: Any,
    parameters: Mapping[str, Param],
    paths: Sequence[PathRow],
    param_defined_names: Mapping[str, str],
    cost_discount_rate: float,
    effect_discount_rate: float,
    variables: Mapping[str, Any] | None,
    distributions: Mapping[str, Mapping[str, Any]] | None,
    tables: Mapping[str, Mapping[str, Any]] | None,
) -> dict[str, dict[str, str]]:
    """One live sheet per PSM node; returns {node id: {total name: cell ref}} for the Tree Paths links."""

    from .psm_excel_export import PSMExportError, psm_sheet_name, write_psm_sheet
    from .psm_trace import compute_psm_trace, find_psm_nodes

    links: dict[str, dict[str, str]] = {}
    used: set[str] = set()
    for node in find_psm_nodes(tree):
        if node.id in links:
            continue
        try:
            trace = compute_psm_trace(
                node, parameters, variables, distributions, tables,
                cost_discount_rate, effect_discount_rate,
            )
        except (ValueError, TypeError) as exc:
            raise PSMExportError(f"PSM node '{node.id}' could not be replayed for the Excel export: {exc}") from exc
        leaf_rows = [row for row in paths if row.leaf_id == node.id and row.leaf_type == "psm"]
        if not leaf_rows:
            continue
        engine_leaf = leaf_rows[0].steps[-1]
        for key, mine, theirs in (
            ("cost", trace.total_cost, engine_leaf.discounted.get("cost", 0.0)),
            ("effect", trace.total_effect, engine_leaf.discounted.get("effect", 0.0)),
            ("cost (undiscounted)", trace.total_cost_undiscounted, engine_leaf.undiscounted.get("cost", 0.0)),
            ("effect (undiscounted)", trace.total_effect_undiscounted, engine_leaf.undiscounted.get("effect", 0.0)),
        ):
            if abs(mine - theirs) > 1e-9 + 1e-9 * abs(theirs):
                raise PSMExportError(
                    f"PSM node '{node.id}': the cycle-by-cycle replay gives {key} {mine!r} but the engine gives "
                    f"{theirs!r}; refusing to export a workbook that would not reconcile."
                )
        sheet = psm_sheet_name(node.id, used)
        links[node.id] = write_psm_sheet(wb, node, trace, param_defined_names, sheet)
    return links


def _write_cover_sheet(
    wb: Any, model_name: str, engine_version: str, config_hash: str, ai_statement: bool, has_psm: bool = False
) -> None:
    ws = wb.create_sheet("Read me", 0)
    lines = [
        model_name or "DECIDE Decision Tree -- transparency export",
        f"Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
        f"DECIDE engine version: {engine_version or 'unknown'}",
        f"Model configuration hash: {config_hash}",
        "",
        "How to read this workbook:",
        "- Parameters: every input's resolved value, distribution and source citation. "
        "A parameter whose value comes from a formula/distribution shows the resolved number here "
        "(Excel cannot represent every DECIDE expression as a native formula) -- see Distribution/Dist args.",
        "- Tree Paths: one row per root-to-leaf path. Amber cells are resolved values (a branch's own "
        "cost/effect, evaluated by DECIDE from its formula); everything else -- joint probability, "
        "path totals, expected contributions -- is a live formula you can trace and recompute.",
        "- Strategy Summary / ICER: fully live formulas built from Tree Paths -- change any Parameters "
        "cell and these recalculate.",
    ]
    if has_psm:
        lines += [
            "- PSM <node> sheets (partitioned survival): the PFS/OS curves, state occupancy, per-cycle cost/effect "
            "and discounting are live formulas; a leaf of Tree Paths that ends at a PSM node takes its payoff "
            "from that sheet's totals, so changing a Parameters cell flows through to Strategy Summary and ICER. "
            "The number of cycles is fixed by the export (one table row per cycle).",
        ]
    if ai_statement:
        lines += ["", "No AI/LLM component was used to compute the results in this file."]
    for i, text in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        if i == 1:
            from openpyxl.styles import Font

            cell.font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120
    from openpyxl.styles import Alignment

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_tree_transparency_workbook(
    tree: Any,
    parameters: Mapping[str, Param],
    paths: Sequence[PathRow],
    *,
    cost_discount_rate: float = 0.0,
    effect_discount_rate: float = 0.0,
    model_name: str = "",
    engine_version: str = "",
    baseline_strategy: str | None = None,
    maximize_effect: bool = True,
    willingness_to_pay: float = 0.0,
    extra_outcome_labels: Mapping[str, str] | None = None,
    variables: Mapping[str, Any] | None = None,
    distributions: Mapping[str, Mapping[str, Any]] | None = None,
    tables: Mapping[str, Mapping[str, Any]] | None = None,
) -> bytes:
    """Build the transparency Excel workbook for a Decision Tree model.

    A tree that contains partitioned-survival (PSM) nodes gets one extra live sheet per PSM node (see
    ``psm_excel_export``); pass the same ``variables``/``distributions``/``tables`` the model runs with so the
    per-cycle replay matches the engine. The build is refused (``PSMExportError``) if a PSM sheet's replay does
    not reconcile with the engine's own value for that node.

    ``tree``/``parameters``/``paths`` are exactly what the app already has
    in hand when it runs the model: ``tree`` is the root ``DecisionNode``,
    ``paths`` is ``tree_paths.enumerate_paths(tree, parameters, ...)``
    (pass the SAME ``cost_discount_rate``/``effect_discount_rate`` to both
    calls -- they must match for the leaf discount formulas built here to
    reconcile with the resolved figures ``enumerate_paths`` computed).
    ``tree`` is walked alongside each path's labels purely to find bare
    parameter references worth making live (see ``_live_ref``); nothing
    here re-evaluates the model. Strategy order is taken from the order
    strategies first appear in ``paths`` (i.e. root branch order).
    ``extra_outcome_labels`` maps an Outcomes-table id (beyond the primary
    "cost"/"effect") to its display label; omit/empty for a model with no
    extra outcomes -- this function does not special-case that, it simply
    has nothing to loop over.
    """

    import openpyxl

    extra_outcome_labels = extra_outcome_labels or {}
    extra_outcome_ids = list(extra_outcome_labels)

    strategy_order: list[str] = []
    for row in paths:
        if row.strategy not in strategy_order:
            strategy_order.append(row.strategy)

    wb = openpyxl.Workbook()
    param_info = _write_parameters_sheet(wb, parameters)
    param_defined_names = {name: info["defined_name"] for name, info in param_info.items()}
    psm_links = _write_psm_sheets(
        wb, tree, parameters, paths, param_defined_names, cost_discount_rate, effect_discount_rate,
        variables, distributions, tables,
    )
    paths_info = _write_tree_paths_sheet(
        wb, tree, paths, param_defined_names, cost_discount_rate, effect_discount_rate,
        extra_outcome_ids, extra_outcome_labels, psm_links,
    )
    summary_info = _write_strategy_summary_sheet(wb, paths_info, strategy_order, extra_outcome_ids, extra_outcome_labels)
    _write_icer_sheet(wb, summary_info, strategy_order, baseline_strategy, maximize_effect, willingness_to_pay)
    config_hash = _config_hash(parameters, paths)
    _write_cover_sheet(wb, model_name, engine_version, config_hash, ai_statement=True, has_psm=bool(psm_links))
    # keep the PSM sheets after ICER (they were created first so Tree Paths could link to them)
    for name in [ws.title for ws in wb.worksheets if ws.title.startswith("PSM ")]:
        wb.move_sheet(name, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(name))
    wb.active = 0

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
