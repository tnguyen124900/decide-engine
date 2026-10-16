"""Batch runs: base case + saved scenarios for many model files at once.

Works on the same pure scenario evaluators the Sensitivity Analysis uses
(sensitivity.evaluate_tree_scenario / evaluate_markov_scenario /
evaluate_microsim_scenario), so a batch row equals what the model's own
Run / scenario table shows. Scenarios are read from
``document["sensitivity_analysis"]["scenarios"]`` ({name: {input: value}}),
the table saved by the Scenario analysis mode of each editor.

Command line (one machine, several cores):

    python -m heor_tree_engine.batch model1.decide folder_of_models -o results.xlsx [--workers 4]
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from concurrent.futures import BrokenExecutor, ProcessPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping, Sequence

BASE_CASE = "Base case"
SUPPORTED_KINDS = ("tree", "markov", "microsim")

RESULT_COLUMNS = (
    "File", "Model type", "Scenario", "Strategy", "Total cost", "Total effect",
    "Incremental cost", "Incremental effect", "ICER", "NMB", "Reference strategy",
    "Optimal at WTP", "WTP",
)


@dataclass
class FileOutcome:
    path: str
    kind: str | None = None
    rows: list[dict[str, Any]] = field(default_factory=list)
    messages: list[str] = field(default_factory=list)  # errors / warnings, one per line

    @property
    def status(self) -> str:
        if not self.rows:
            return "Failed"
        return "OK with errors" if self.messages else "OK"


def detect_kind(document: Mapping[str, Any]) -> str | None:
    """tree / markov / microsim from the file content (None when unsupported,
    e.g. NMA, PSM-only docs run as trees since a PSM model is a tree)."""

    if not isinstance(document, Mapping):
        return None
    if "studies" in document or "analyses" in document or "treatments" in document:
        return None  # NMA
    if "n_patients" in document and "strategies" in document:
        return "microsim"
    if "strategies" in document and "cycles" in document:
        return "markov"
    if "root_id" in document and "nodes" in document:
        return "tree"
    return None


def _evaluator(kind: str):
    from . import sensitivity

    return {
        "tree": sensitivity.evaluate_tree_scenario,
        "markov": sensitivity.evaluate_markov_scenario,
        "microsim": sensitivity.evaluate_microsim_scenario,
    }[kind]


def run_document(
    document: Mapping[str, Any],
    file_label: str = "",
    kind: str | None = None,
    scenarios: Mapping[str, Mapping[str, float]] | None = None,
    include_saved_scenarios: bool = True,
) -> FileOutcome:
    """Base case plus the model's saved scenarios, as flat result rows."""

    outcome = FileOutcome(path=file_label)
    kind = kind or detect_kind(document)
    outcome.kind = kind
    if kind not in SUPPORTED_KINDS:
        outcome.messages.append("Unsupported model type for batch run (supported: Tree/PSM, Markov, Microsimulation).")
        return outcome
    config = document.get("sensitivity_analysis") or {}
    wtp = float(config.get("willingness_to_pay", 50_000.0) or 50_000.0)
    baseline = (document.get("analysis") or {}).get("baseline_strategy_id") or None
    evaluate = _evaluator(kind)

    to_run: dict[str, Mapping[str, float]] = {BASE_CASE: {}}
    if scenarios is None and include_saved_scenarios:
        scenarios = config.get("scenarios") or {}
    for name, overrides in (scenarios or {}).items():
        if str(name).strip().lower() == BASE_CASE.lower() and not overrides:
            continue  # the base case is always included once
        to_run[str(name)] = overrides or {}

    for name, overrides in to_run.items():
        try:
            result = evaluate(document, overrides, baseline, wtp)
        except Exception as exc:  # noqa: BLE001 -- one failing scenario must not stop the batch
            outcome.messages.append(f"{name}: {exc}")
            continue
        if result.errors:
            outcome.messages.append(f"{name}: " + "; ".join(str(e) for e in result.errors))
            continue
        for item in result.outcomes:
            outcome.rows.append({
                "File": file_label, "Model type": kind, "Scenario": name,
                "Strategy": item.strategy_name or item.strategy_id,
                "Total cost": item.total_cost, "Total effect": item.total_effect,
                "Incremental cost": item.incremental_cost, "Incremental effect": item.incremental_effect,
                "ICER": item.icer, "NMB": item.nmb,
                "Reference strategy": result.reference_strategy_id,
                "Optimal at WTP": item.strategy_id == result.optimal_strategy_id,
                "WTP": result.willingness_to_pay,
            })
    return outcome


def run_file(path: str) -> FileOutcome:
    """Load one model file and run it. Never raises."""

    label = os.path.basename(path)
    try:
        with open(path, encoding="utf-8") as handle:
            document = json.load(handle)
    except (OSError, ValueError) as exc:
        return FileOutcome(path=label, messages=[f"Could not read the file: {exc}"])
    try:
        return run_document(document, file_label=label)
    except Exception as exc:  # noqa: BLE001
        return FileOutcome(path=label, messages=[f"Run failed: {exc}"])


def collect_files(inputs: Sequence[str]) -> list[str]:
    """Files given directly, plus every .decide/.json file inside folders."""

    found: list[str] = []
    for item in inputs:
        p = Path(item)
        if p.is_dir():
            found.extend(str(f) for f in sorted(list(p.glob("*.decide")) + list(p.glob("*.json"))))
        elif p.is_file():
            found.append(str(p))
    seen: set[str] = set()
    return [f for f in found if not (f in seen or seen.add(f))]


def run_batch(paths: Sequence[str], workers: int | None = None) -> list[FileOutcome]:
    """Run every file (one process per file when ``workers`` > 1). Order of
    the returned outcomes follows ``paths``."""

    paths = list(paths)
    if not paths:
        return []
    workers = workers or min(os.cpu_count() or 1, len(paths))
    if workers <= 1 or len(paths) == 1:
        return [run_file(p) for p in paths]
    try:
        with ProcessPoolExecutor(max_workers=workers) as pool:
            return list(pool.map(run_file, paths))
    except (RuntimeError, OSError, AssertionError, BrokenExecutor):
        # No sub-processes allowed here (e.g. inside a daemon process of a
        # packaged app): run one after the other instead.
        return [run_file(p) for p in paths]


def write_excel(outcomes: Sequence[FileOutcome], path: Any) -> None:
    """Results sheet (one row per file x scenario x strategy) + Files sheet
    (status and messages per file). ``path`` may be a file name or a
    file-like object (e.g. BytesIO)."""

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(list(RESULT_COLUMNS))
    for outcome in outcomes:
        for row in outcome.rows:
            ws.append([row.get(c) for c in RESULT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, column in enumerate(RESULT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(column) + 2)
        if column in ("Total cost", "Incremental cost", "NMB"):
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "#,##0.00"
        elif column in ("Total effect", "Incremental effect", "ICER"):
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "#,##0.0000"

    files = wb.create_sheet("Files")
    files.append(["File", "Model type", "Status", "Rows", "Messages"])
    for cell in files[1]:
        cell.font = Font(bold=True)
    for outcome in outcomes:
        files.append([outcome.path, outcome.kind or "", outcome.status, len(outcome.rows), "\n".join(outcome.messages)])
    files.column_dimensions["A"].width = 36
    files.column_dimensions["E"].width = 90
    wb.save(path)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run base case + saved scenarios for many DECIDE model files.")
    parser.add_argument("inputs", nargs="+", help="model files and/or folders of .decide/.json files")
    parser.add_argument("-o", "--out", default="decide-batch-results.xlsx", help="Excel file to write")
    parser.add_argument("--workers", type=int, default=None, help="processes (default: one per core, max one per file)")
    args = parser.parse_args(argv)
    files = collect_files(args.inputs)
    if not files:
        print("No model files found.", file=sys.stderr)
        return 2
    outcomes = run_batch(files, args.workers)
    write_excel(outcomes, args.out)
    for outcome in outcomes:
        print(f"{outcome.status:15s} {outcome.path}")
        for message in outcome.messages:
            print(f"    {message}")
    print(f"Wrote {args.out}")
    return 0 if any(o.rows for o in outcomes) else 1


if __name__ == "__main__":
    raise SystemExit(main())
